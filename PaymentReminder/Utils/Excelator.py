from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

class ExcelManager:

    registros = [
        {
            "Identificador": 1,
            "Talento": "Nombre del talento",
            "Marca": "Nombre de la marca",
            "Cantidad Total": 2500.00,
            "Fecha de Facturación": "2026-09-30",
            "Días Restantes": 15,
            "Cobrado": False,
        },
    ]

    columns = [
        "Identificador",
        "Talento",
        "Marca",
        "Cantidad Total",
        "Fecha de Facturación",
        "Días Restantes",
        "Cobrado",
    ]

    @staticmethod
    def generateExcel(registros=registros, columns=columns, nombre_archivo="payments.xlsx"):
        libro = Workbook()
        hoja = libro.active
        hoja.title = "Payments"

        hoja.append(columns)

        for registro in registros:
            hoja.append([registro.get(columna) for columna in columns])

        for celda in hoja[1]:
            celda.font = Font(bold=True)

        hoja.freeze_panes = "A2"
        hoja.auto_filter.ref = hoja.dimensions

        for columna in hoja.columns:
            ancho = max(len(str(celda.value or "")) for celda in columna) + 2
            letra = get_column_letter(columna[0].column)
            hoja.column_dimensions[letra].width = ancho

        libro.save(nombre_archivo)